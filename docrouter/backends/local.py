"""Pure-Python backends: the floor and the cheap-but-competent tier.

`naive` exists to be beaten. It is the "just rip the text out" strawman that
most ad-hoc pipelines actually ship, and having it in the benchmark keeps the
other backends honest about how much structure they are really recovering.

`pylib` is the free CPU tier: no models, no GPU, no network. On clean digital
documents it is often within noise of the expensive backends, which is exactly
the finding that makes a cost-aware router worth building.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from .base import Backend

DOCX = {".docx"}
PDF = {".pdf"}


def _clean(text: str) -> str:
    return " ".join(text.split())


class NaiveBackend(Backend):
    """Strip the markup, keep whatever falls out. Structure is destroyed."""

    name = "naive"
    cost_per_page_usd = 0.0

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in DOCX | PDF

    def _parse(self, path: Path) -> tuple[str, int, dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in DOCX:
            with zipfile.ZipFile(path) as z:
                xml = z.read("word/document.xml").decode("utf-8")
            text = _clean(re.sub(r"<[^>]+>", " ", xml))
            return text, 1, {"strategy": "xml-strip"}

        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(path) as pdf:
            pages = len(pdf.pages)
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        # Deliberately no page separators and no layout mode: this is the
        # glyph dump that interleaves multi-column pages.
        return _clean(" ".join(parts)), pages, {"strategy": "glyph-dump"}


class PyLibBackend(Backend):
    """python-docx / pdfplumber with structure preserved as GFM markdown."""

    name = "pylib"
    cost_per_page_usd = 0.0

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in DOCX | PDF

    def is_available(self) -> bool:
        try:
            import docx  # noqa: F401
            import pdfplumber  # noqa: F401
        except ImportError:
            return False
        return True

    def _parse(self, path: Path) -> tuple[str, int, dict[str, Any]]:
        if path.suffix.lower() in DOCX:
            return self._parse_docx(path)
        return self._parse_pdf(path)

    def _parse_docx(self, path: Path) -> tuple[str, int, dict[str, Any]]:
        import docx

        doc = docx.Document(str(path))
        blocks: list[str] = []
        n_tables = 0

        # python-docx exposes paragraphs and tables as separate collections, so
        # walk the body XML directly to keep them in document order. Getting
        # this wrong is a silent reading-order bug that only shows up in the
        # reading-order metric, which is a good argument for having that metric.
        body = doc.element.body
        para_iter = iter(doc.paragraphs)
        table_iter = iter(doc.tables)
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = next(para_iter, None)
                if para is None or not para.text.strip():
                    continue
                blocks.append(_docx_para_to_md(para))
            elif tag == "tbl":
                table = next(table_iter, None)
                if table is None:
                    continue
                n_tables += 1
                blocks.append(_docx_table_to_md(table))

        return "\n\n".join(blocks), 1, {"tables": n_tables}

    def _parse_pdf(self, path: Path) -> tuple[str, int, dict[str, Any]]:
        import pdfplumber

        blocks: list[str] = []
        n_tables = 0
        n_fragments = 0
        with pdfplumber.open(path) as pdf:
            pages = len(pdf.pages)
            for page in pdf.pages:
                raw_tables = page.find_tables()
                n_fragments += len(raw_tables)
                groups = _group_table_fragments(raw_tables)
                n_tables += len(groups)
                table_boxes = [
                    _group_bbox(g) for g in groups
                ]

                # Remove table regions before pulling body text, otherwise every
                # cell value appears twice: once as prose, once in the table.
                body = page
                for bbox in table_boxes:
                    body = body.filter(
                        lambda obj, bbox=bbox: not _inside(obj, bbox)
                    )

                # Do NOT use extract_text(layout=True) here: it preserves
                # spatial position, which splices two-column pages together
                # character by character. Reconstruct reading order instead.
                from ..layout import lines_to_paragraphs, order_words

                words = body.extract_words()
                lines = order_words(words, page.width)
                blocks.extend(lines_to_paragraphs(lines))

                for group in groups:
                    rows: list[list[str | None]] = []
                    for table in group:
                        rows.extend(table.extract() or [])
                    if len(rows) >= 2:
                        blocks.append(_rows_to_md(rows))

        return (
            "\n\n".join(blocks),
            pages,
            {"tables": n_tables, "raw_fragments": n_fragments},
        )


# Tolerances for stitching table fragments back together.
X_TOLERANCE = 3.0   # points; fragments of one table share left/right edges
GAP_FACTOR = 2.5    # multiples of fragment height allowed between fragments


def _group_bbox(group: list[Any]) -> tuple[float, float, float, float]:
    return (
        min(t.bbox[0] for t in group),
        min(t.bbox[1] for t in group),
        max(t.bbox[2] for t in group),
        max(t.bbox[3] for t in group),
    )


def _group_table_fragments(tables: list[Any]) -> list[list[Any]]:
    """Stitch single-row table fragments back into whole tables.

    pdfplumber's line-based detection returns one Table per ruled row on some
    filings: a 97-page Microsoft 10-Q produced 235 "tables", most of them a
    single row like ['Revenue:', '', '', ...]. Downstream every one is
    discarded, because a table needs at least two rows to be a table. The
    document scored TEDS 0.011 against 41 real tables -- not because structure
    was recovered badly, but because it was never assembled.

    Fragments of one table share left and right edges almost exactly and stack
    with roughly one line of vertical gap. Group on that, concatenate the rows,
    and the table comes back.

    Tables that were detected correctly in the first place are single-element
    groups, so this is a no-op on documents that never had the problem.
    """
    if not tables:
        return []

    ordered = sorted(tables, key=lambda t: (t.bbox[1], t.bbox[0]))
    groups: list[list[Any]] = [[ordered[0]]]

    for table in ordered[1:]:
        prev = groups[-1][-1]
        px0, _, px1, pbottom = prev.bbox
        x0, top, x1, bottom = table.bbox

        same_columns = (
            abs(x0 - px0) <= X_TOLERANCE and abs(x1 - px1) <= X_TOLERANCE
        )
        height = max(bottom - top, pbottom - prev.bbox[1], 1.0)
        adjacent = 0 <= (top - pbottom) <= height * GAP_FACTOR

        if same_columns and adjacent:
            groups[-1].append(table)
        else:
            groups.append([table])

    return groups


def _inside(obj: dict[str, Any], bbox: tuple[float, float, float, float]) -> bool:
    x0, top, x1, bottom = bbox
    cx = (obj["x0"] + obj["x1"]) / 2
    cy = (obj["top"] + obj["bottom"]) / 2
    return x0 <= cx <= x1 and top <= cy <= bottom


def _clean_layout_text(text: str) -> str:
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line.strip())


def _docx_para_to_md(para: Any) -> str:
    style = (para.style.name or "").lower()
    text = para.text.strip()
    if style.startswith("heading"):
        digits = "".join(c for c in style if c.isdigit())
        level = int(digits) if digits else 1
        return f"{'#' * min(level, 6)} {text}"
    if style.startswith("list"):
        return f"- {text}"
    return text


def _docx_table_to_md(table: Any) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    return _rows_to_md(rows)


def _rows_to_md(rows: list[list[str | None]]) -> str:
    """Render a row matrix as a GitHub-flavored pipe table."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [
        [(c or "").replace("\n", " ").replace("|", r"\|").strip() for c in r]
        + [""] * (width - len(r))
        for r in rows
    ]
    header, *body = norm
    out = ["| " + " | ".join(header) + " |"]
    out.append("| " + " | ".join(["---"] * width) + " |")
    for row in body:
        out.append("| " + " | ".join(row) + " |")
    return "\n".join(out)
