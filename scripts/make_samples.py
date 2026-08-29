"""Generate a small synthetic corpus with known ground truth.

This is a smoke-test corpus, not a benchmark. Its job is to prove the harness
runs and that the metrics move in the direction they should: the naive backend
must score badly on the table document and badly on the two-column PDF, and if
it does not, the metric is broken rather than the backend being good.

Replace this with real annotated documents before reporting any number you
intend to put on a resume.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent / "data" / "sample_corpus"
DOCS = ROOT / "docs"
GOLD = ROOT / "gold"

TABLE_ROWS = [
    ("Region", "Revenue ($M)", "Change vs Q2"),
    ("North", "128.4", "+6.1%"),
    ("South", "95.2", "+2.8%"),
    ("East", "143.9", "+11.4%"),
    ("West", "61.7", "-4.3%"),
]

INTRO = (
    "Revenue grew in three of four regions. The board flagged the West region "
    "for review after its second consecutive quarterly decline."
)

LEFT_TEXT = (
    "The committee reviewed operating expenses across all four regions during "
    "the quarter and found that headcount growth accounted for the majority of "
    "the increase. Facilities costs remained flat year over year despite the "
    "opening of two satellite offices in the eastern territory."
)

RIGHT_TEXT = (
    "Marketing spend rose sharply in the second half of the period, driven by "
    "the product launch campaign. The committee recommends holding the current "
    "allocation through the next quarter and revisiting the mix once attribution "
    "data from the campaign becomes available for analysis."
)


def wrap_to_width(text: str, max_width: float, font: str, size: float) -> list[str]:
    """Greedy wrap using real glyph widths.

    Wrapping by character count is what produced the first broken version of
    this file: the left column overran the gutter and physically overlapped the
    right column, so the "two-column PDF" was not two columns at all. Measure
    the font.
    """
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines: list[str] = []
    current: list[str] = []
    for word in text.split():
        trial = " ".join(current + [word])
        if current and stringWidth(trial, font, size) > max_width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return lines


def make_table_docx() -> str:
    doc = Document()
    doc.add_heading("Q3 Financial Summary", level=1)
    doc.add_paragraph(INTRO)
    table = doc.add_table(rows=len(TABLE_ROWS), cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(TABLE_ROWS):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
    doc.save(DOCS / "0001_table.docx")

    header, *body = TABLE_ROWS
    md = [
        "# Q3 Financial Summary",
        "",
        INTRO,
        "",
        "| " + " | ".join(header) + " |",
        "| --- | --- | --- |",
    ]
    md += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(md) + "\n"


def make_two_column_pdf() -> str:
    path = DOCS / "0002_twocol.pdf"
    c = canvas.Canvas(str(path), pagesize=LETTER)
    width, height = LETTER

    # Two 3-inch columns with a half-inch gutter between them. The gutter has
    # to be genuinely empty or there is nothing for a layout model to find.
    col_width = 3 * inch
    left_x = inch
    right_x = left_x + col_width + 0.5 * inch

    c.setFont("Helvetica-Bold", 14)
    c.drawString(left_x, height - inch, "Operating Expense Review")

    c.setFont("Helvetica", 10)
    for x, text in ((left_x, LEFT_TEXT), (right_x, RIGHT_TEXT)):
        y = height - 1.5 * inch
        for line in wrap_to_width(text, col_width, "Helvetica", 10):
            c.drawString(x, y, line)
            y -= 14

    c.save()

    return "\n\n".join([
        "# Operating Expense Review",
        LEFT_TEXT,
        RIGHT_TEXT,
    ]) + "\n"


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    GOLD.mkdir(parents=True, exist_ok=True)

    gold = {
        "0001_table": make_table_docx(),
        "0002_twocol": make_two_column_pdf(),
    }
    for key, markdown in gold.items():
        (GOLD / f"{key}.md").write_text(markdown, encoding="utf-8")

    meta = [
        {"doc_key": "0001_table", "columns": 1, "scanned": False,
         "table_density": "high", "source": "synthetic"},
        {"doc_key": "0002_twocol", "columns": 2, "scanned": False,
         "table_density": "none", "source": "synthetic"},
    ]
    with open(ROOT / "meta.jsonl", "w", encoding="utf-8") as fh:
        for record in meta:
            fh.write(json.dumps(record) + "\n")

    print(f"Wrote {len(gold)} documents to {ROOT}")


if __name__ == "__main__":
    main()
